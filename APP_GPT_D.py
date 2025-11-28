import os
import shutil
import asyncio
import re
import PyPDF2
import docx
import pandas as pd
from typing import List, Tuple, Any, Dict, Optional
import uvicorn

from fastapi import FastAPI, Request, UploadFile, File
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

# --- IMPORTS LANGCHAIN ---
from langchain_text_splitters import RecursiveCharacterTextSplitter
# Import du Semantic Chunker (nécessite langchain-experimental)
from langchain_experimental.text_splitter import SemanticChunker

from langchain_openai import OpenAIEmbeddings
from langchain_community.vectorstores import FAISS
from langchain.chains import ConversationalRetrievalChain
from langchain_core.documents import Document
from langchain_core.prompts import PromptTemplate
from langchain_openai import ChatOpenAI

# Import de l'Agent Excel (nécessite langchain-experimental)
from langchain_experimental.agents.agent_toolkits import create_pandas_dataframe_agent

from langdetect import detect, LangDetectException
from dotenv import load_dotenv

load_dotenv()

# ----------------------------
# Configuration
# ----------------------------
# Gestion de la persistance (Volume Railway vs Local)
BASE_DIR = "/app/storage" if os.path.exists("/app/storage") else "."

FOLDER = os.path.join(BASE_DIR, "Docu_Cost_notTechnip")
VECTOR_INDEX_PATH = os.path.join(BASE_DIR, "vector_index_cost_GPT")

RETRIEVER_K = 6
MAX_SOURCES_RETURN = 8

OPENAI_CHAT_MODEL = "gpt-5-nano-2025-08-07"
OPENAI_EMBEDDING_MODEL = "text-embedding-3-large"

MARKER_FR = "Le contexte fourni n'a pas de rapport avec cette question."
MARKER_EN = "The provided context is not relevant to this question."

# ----------------------------
# PROMPTS
# ----------------------------

# 1. FRANÇAIS
raw_template_fr = """Tu es un assistant IA. Réponds à la "Question" en te basant sur l'historique de la conversation et le "Contexte" fourni.
Règle cruciale : La réponse doit impérativement être en français.

MARQUEUR SPÉCIAL (À UTILISER TEL QUEL) :
Si le "Contexte" n'est pas pertinent, commence ta réponse par la phrase exacte : "MARKER_PLACEHOLDER" (avec l'espace à la fin).

INSTRUCTIONS :
1.  Priorité 1 : Utilise le "Contexte" ci-dessous pour répondre de manière factuelle à la "Question".
2.  Priorité 2 : Si le "Contexte" ne contient pas la réponse ou n'est pas pertinent, utilise le "MARQUEUR SPÉCIAL" puis réponds à la "Question" en utilisant tes connaissances générales et l'historique de la conversation.
3.  Priorité 3 : Si tu ne peux pas répondre du tout, dis-le poliment. N'invente pas une réponse.

Contexte:
{context}

Question:
{question}

Réponse:"""

final_template_fr = raw_template_fr.replace("MARKER_PLACEHOLDER", MARKER_FR)
PROMPT_FR = PromptTemplate(template=final_template_fr, input_variables=["context", "question"])

# 2. ANGLAIS
raw_template_en = """You are an AI assistant. Answer the "Question" based on the conversation history and the provided "Context".
Crucial rule: The answer must absolutely be in English.

SPECIAL MARKER (USE AS IS):
If the "Context" is not relevant, start your answer with the exact phrase: "MARKER_PLACEHOLDER" (with the space at the end).

INSTRUCTIONS:
1.  Priority 1: Use the "Context" below to factually answer the "Question".
2.  Priority 2: If the "Context" does not contain the answer or is not relevant, use the "SPECIAL MARKER" and then answer the "Question" using your general knowledge and the conversation history.
3.  Priority 3: If you cannot answer at all, state so politely. Do not invent an answer.

Context:
{context}

Question:
{question}

Answer:"""

final_template_en = raw_template_en.replace("MARKER_PLACEHOLDER", MARKER_EN)
PROMPT_EN = PromptTemplate(template=final_template_en, input_variables=["context", "question"])

# 3. CONDENSATION
CONDENSE_QUESTION_PROMPT_TEMPLATE = """Étant donné l'historique de la conversation (Chat History) et une question de suivi (Follow Up Input), reformule la question de suivi pour qu'elle soit une question autonome.
Historique: {chat_history}
Question de suivi: {question}
Question autonome:"""
CONDENSE_PROMPT = PromptTemplate.from_template(CONDENSE_QUESTION_PROMPT_TEMPLATE)

# ----------------------------
# SETUP APP
# ----------------------------
app = FastAPI()

if os.path.exists("static"):
    app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    return FileResponse("static/images/logo.png")

class ChatRequest(BaseModel):
    query: str
    chat_history: List[Any] = []
    session_id: Optional[str] = "global"

# Variables Globales
chatbot_instances: Dict[str, ConversationalRetrievalChain] = {}
vectorstore_global: Optional[FAISS] = None
excels_global: Dict[str, Dict[str, pd.DataFrame]] = {}
locks_by_session: Dict[str, asyncio.Lock] = {}

# ----------------------------
# HELPERS (Extraction & Loading)
# ----------------------------
def extract_docs_from_pdf(path: str) -> List[Document]:
    docs: List[Document] = []
    try:
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            if getattr(reader, "is_encrypted", False):
                try: reader.decrypt("")
                except: return docs
            for i, page in enumerate(reader.pages, start=1):
                txt = page.extract_text() or ""
                if txt.strip():
                    docs.append(Document(page_content=txt, metadata={"source": os.path.basename(path), "page": i}))
    except Exception as e:
        print(f"⚠️ Erreur lecture PDF {path}: {e}")
    return docs

def extract_docs_from_docx(path: str) -> List[Document]:
    docs: List[Document] = []
    try:
        d = docx.Document(path)
        paragraphs = [p.text for p in d.paragraphs if p.text and p.text.strip()]
        if paragraphs:
            docs.append(Document(page_content="\n".join(paragraphs), metadata={"source": os.path.basename(path), "page": 1}))
    except Exception as e:
        print(f"⚠️ Erreur lecture DOCX {path}: {e}")
    return docs

def load_documents(folder: str) -> List[Document]:
    all_docs: List[Document] = []
    if not os.path.exists(folder): 
        print(f"ℹ️ Le dossier {folder} n'existe pas encore.")
        return all_docs
    files = os.listdir(folder)
    print(f"📂 Fichiers trouvés: {files}")
    valid_files = [f for f in files if f.lower().endswith((".pdf", ".docx")) and not f.startswith("~$")]
    for file in valid_files:
        full_path = os.path.join(folder, file)
        docs = []
        try:
            if file.lower().endswith(".pdf"): docs = extract_docs_from_pdf(full_path)
            elif file.lower().endswith(".docx"): docs = extract_docs_from_docx(full_path)
            if docs: all_docs.extend(docs)
        except Exception as e:
            print(f"❌ CRASH lecture {file}: {e}")
    return all_docs

def load_excel(path: str) -> Dict[str, pd.DataFrame]:
    try:
        xls = pd.ExcelFile(path)
        return {sheet: pd.read_excel(path, sheet_name=sheet) for sheet in xls.sheet_names}
    except Exception as e:
        print(f"⚠️ Erreur lecture Excel {path} : {e}")
        return {}

def detect_language_smart(query: str) -> str:
    """Détection langue avec correction pour phrases courtes."""
    query_clean = query.strip().lower()
    french_keywords = ["oui", "non", "vas-y", "vas y", "merci", "bonjour", "salut", "d'accord", "ok", "c'est", "ça", "plaît", "stp", "continue", "encore"]
    if len(query_clean) < 50:
        if any(word in query_clean for word in french_keywords): return 'fr'
    try: return detect(query)
    except LangDetectException: return 'en'

# --- FONCTIONS MANQUANTES AJOUTÉES ICI ---

def build_chat_history_for_chain(chat_history_raw: List[Any], max_turns: int = 6) -> List[Tuple[str, str]]:
    if not chat_history_raw: return []
    pairs: List[Tuple[str, str]] = []
    last_user = None
    for item in chat_history_raw:
        if isinstance(item, dict):
            role = item.get("role", "").lower()
            content = item.get("content", "")
            if role == "user": last_user = content
            elif role == "assistant" and last_user is not None:
                pairs.append((last_user, content))
                last_user = None
    return pairs[-max_turns:]

def extract_unique_sources(result_obj: Dict[str, Any], max_sources: int) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    seen = set()
    docs = result_obj.get("source_documents", [])
    if not docs: return out
    for d in docs:
        meta = d.metadata
        src_norm = meta.get("source", "inconnu")
        page = meta.get("page", "?")
        key = (src_norm, page)
        if key not in seen:
            seen.add(key)
            out.append({"Document": src_norm, "Page/Feuille": str(page)})
            if len(out) >= max_sources: break
    return out

# ----------------------------
# LOGIQUE EXCEL (AGENT PANDAS)
# ----------------------------
def analyse_excel(query: str, excels: Dict[str, Dict[str, pd.DataFrame]]) -> Tuple[Optional[str], Optional[List[Dict[str, str]]]]:
    # Mots-clés pour déclencher l'analyseur mathématique
    triggers = ["excel", "tableau", "combien", "somme", "total", "moyenne", "coût", "prix", "max", "min", "colonne", "ligne", "donnée", "chiffre", "data"]
    if not any(t in query.lower() for t in triggers):
        return None, None
    
    # Aplatir les données
    df_list = []
    for filename, sheets in excels.items():
        for sheetname, df in sheets.items():
            if not df.empty: df_list.append(df)
    
    if not df_list: return None, None
    
    try:
        # Création de l'agent
        llm_excel = ChatOpenAI(model=OPENAI_CHAT_MODEL, temperature=0)
        agent = create_pandas_dataframe_agent(
            llm_excel, 
            df_list, 
            verbose=True, 
            allow_dangerous_code=True, 
            handle_parsing_errors=True
        )
        print(f"📊 Interrogation Agent Excel pour : {query}")
        response = agent.invoke(query)
        result_text = response.get("output", "Je n'ai pas pu calculer la réponse.")
        return result_text, [{"Document": "Analyse Excel", "Page/Feuille": "Calcul IA"}]
    except Exception as e:
        print(f"⚠️ Erreur Agent Excel : {e}")
        return None, None

# ----------------------------
# LOGIQUE RAG (SEMANTIC CHUNKING)
# ----------------------------
def get_or_create_vector_db(docs: List[Document], index_path: str = VECTOR_INDEX_PATH):
    embeddings = OpenAIEmbeddings(model=OPENAI_EMBEDDING_MODEL)
    
    # 1. Chargement persistant si dispo
    if os.path.exists(index_path) and os.path.exists(os.path.join(index_path, "index.faiss")):
        try: 
            print("📂 Chargement de l'index existant...")
            return FAISS.load_local(index_path, embeddings, allow_dangerous_deserialization=True)
        except: pass
        
    if not docs: return None

    # 2. CONSTRUCTION AVEC SEMANTIC CHUNKING
    print("🧠 Démarrage du Semantic Chunking...")
    splitter = SemanticChunker(embeddings, breakpoint_threshold_type="percentile")
    
    try:
        chunks = splitter.split_documents(docs)
        print(f"✂️  Documents découpés en {len(chunks)} blocs sémantiques.")
    except Exception as e:
        print(f"⚠️ Fallback découpage classique ({e})")
        splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=150)
        chunks = splitter.split_documents(docs)

    vectorstore = FAISS.from_documents(chunks, embeddings)
    
    # 3. Sauvegarde
    try: 
        if not os.path.exists(os.path.dirname(index_path)): os.makedirs(os.path.dirname(index_path), exist_ok=True)
        vectorstore.save_local(index_path)
    except: pass
    return vectorstore

def create_chatbot(vectorstore: FAISS, prompt: PromptTemplate):
    if not vectorstore: return None
    llm = ChatOpenAI(model=OPENAI_CHAT_MODEL, temperature=1)
    retriever = vectorstore.as_retriever(search_kwargs={"k": RETRIEVER_K})
    return ConversationalRetrievalChain.from_llm(
        llm=llm, retriever=retriever, return_source_documents=True,
        combine_docs_chain_kwargs={"prompt": prompt}, condense_question_prompt=CONDENSE_PROMPT
    )

def build_all_chatbots():
    global chatbot_instances, vectorstore_global
    if not os.getenv("OPENAI_API_KEY"): return
    docs = load_documents(FOLDER)
    vectorstore_global = get_or_create_vector_db(docs)
    if vectorstore_global:
        chatbot_instances['fr'] = create_chatbot(vectorstore_global, PROMPT_FR)
        chatbot_instances['en'] = create_chatbot(vectorstore_global, PROMPT_EN)

@app.on_event("startup")
async def startup_event():
    # Création dossiers
    if not os.path.exists(FOLDER): 
        try: os.makedirs(FOLDER, exist_ok=True)
        except: pass
    
    # Chargement Excel
    global excels_global
    if os.path.exists(FOLDER):
        excel_files = [f for f in os.listdir(FOLDER) if f.endswith((".xlsx", ".xls")) and not f.startswith("~$")]
        for file in excel_files:
            excels_global[file] = load_excel(os.path.join(FOLDER, file))
            
    build_all_chatbots()

# ----------------------------
# ROUTES
# ----------------------------
@app.get("/", response_class=HTMLResponse)
async def root():
    if os.path.exists("index.html"):
        with open("index.html", "r", encoding="utf-8") as f: return HTMLResponse(f.read())
    return HTMLResponse("<h3>API RAG & Excel Ready (Semantic)</h3>")

@app.post("/api/upload")
async def upload_document(file: UploadFile = File(...)):
    global vectorstore_global, excels_global
    filename = file.filename
    file_path = os.path.join(FOLDER, filename)
    try:
        with open(file_path, "wb") as buffer: shutil.copyfileobj(file.file, buffer)
    except Exception as e: return JSONResponse({"status": "error", "message": str(e)})

    # EXCEL
    if filename.lower().endswith((".xlsx", ".xls")):
        excels_global[filename] = load_excel(file_path)
        return JSONResponse({"status": "success", "message": f"Excel '{filename}' chargé."})

    # RAG (PDF/DOCX)
    new_docs = []
    if filename.lower().endswith(".pdf"): new_docs = extract_docs_from_pdf(file_path)
    elif filename.lower().endswith(".docx"): new_docs = extract_docs_from_docx(file_path)
    
    if not new_docs: return JSONResponse({"status": "error", "message": "Aucun texte valide."})

    if vectorstore_global:
        try:
            # Ajout dynamique avec Semantic Chunking
            embeddings = OpenAIEmbeddings(model=OPENAI_EMBEDDING_MODEL)
            splitter = SemanticChunker(embeddings, breakpoint_threshold_type="percentile")
            chunks = splitter.split_documents(new_docs)
            
            vectorstore_global.add_documents(chunks)
            vectorstore_global.save_local(VECTOR_INDEX_PATH)
            return JSONResponse({"status": "success", "message": f"Fichier '{filename}' ajouté (Semantic Chunking)."})
        except Exception as e: return JSONResponse({"status": "error", "message": str(e)})
    else:
        try:
            vectorstore_global = get_or_create_vector_db(new_docs)
            chatbot_instances['fr'] = create_chatbot(vectorstore_global, PROMPT_FR)
            chatbot_instances['en'] = create_chatbot(vectorstore_global, PROMPT_EN)
            return JSONResponse({"status": "success", "message": f"Base initialisée avec '{filename}'."})
        except Exception as e: return JSONResponse({"status": "error", "message": str(e)})

@app.post("/api/chat")
async def chat_endpoint(request: ChatRequest):
    global chatbot_instances, excels_global, locks_by_session
    query = (request.query or "").strip()
    if not query: return JSONResponse({"answer": "Question vide.", "sources": []})
    session_id = request.session_id or "global"
    lock = locks_by_session.setdefault(session_id, asyncio.Lock())

    async with lock:
        # 1. Analyse Excel
        excel_answer, excel_sources = analyse_excel(query, excels_global)
        if excel_answer: return JSONResponse({"answer": excel_answer, "sources": excel_sources})
        
        # 2. RAG Classique
        lang = detect_language_smart(query)
        chatbot_instance = chatbot_instances.get(lang, chatbot_instances.get('en'))
        if not chatbot_instance: return JSONResponse({"answer": "⚠️ Chatbot non prêt.", "sources": []})
        
        # Utilisation des fonctions utilitaires
        chat_history_tuples = build_chat_history_for_chain(request.chat_history, max_turns=6)
        try:
            result = await asyncio.to_thread(chatbot_instance.invoke, {"question": query, "chat_history": chat_history_tuples})
        except Exception as e:
            print(f"ERREUR CHAT: {e}")
            return JSONResponse({"answer": "⚠️ Erreur interne.", "sources": []})

        answer = result.get("answer", "")
        sources = extract_unique_sources(result, max_sources=MAX_SOURCES_RETURN)
        
        if answer.strip().startswith(MARKER_FR) or answer.strip().startswith(MARKER_EN): sources = []
        return JSONResponse({"answer": answer, "sources": sources})

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("APP_GPT_D:app", host="0.0.0.0", port=port)